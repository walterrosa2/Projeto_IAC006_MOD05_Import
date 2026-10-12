"""Testes de integração da esteira: pasta -> classificação -> IA -> aprovação -> exportação.

Cobrem os critérios de aceite CA01 a CA08 do PRD.
"""

from __future__ import annotations

from openpyxl import load_workbook

from core.accounting.entry_generator import generate_entry_for_document
from core.constants import DocumentStatus, DocumentType, EntryStatus
from core.dominio.dominio_exporter import export_entries
from core.pipeline import process_company_folder
from core.scanner.file_scanner import scan_folder
from db.repositories import accounting_entry_repository, document_repository

CONTAS_PAGAS = """RELATORIO DE CONTAS PAGAS
Empresa: CLIENTE A LTDA
CREDOR: FORNECEDOR FICTICIO SERVICOS LTDA
DOCUMENTO: NF-2026-0042
DATA DE BAIXA: 06/01/2026
VALOR BRUTO: 1.600,00
DESCONTO: 100,00
VALOR LIQUIDO: 1.500,00
"""

PIX = """COMPROVANTE DE TRANSFERENCIA PIX
Data da transferência: 06/01/2026 14:35
Valor: R$ 1.500,00
PAGADOR
CLIENTE A LTDA
FAVORECIDO
FORNECEDOR FICTICIO SERVICOS LTDA
ID: E12345678202601061435
"""

ILEGIVEL = "conteúdo sem qualquer relação com documentos contábeis conhecidos"


def _write(company, name: str, content: str):
    from pathlib import Path

    path = Path(company.input_folder_path) / name
    path.write_text(content, encoding="utf-8")
    return path


class TestCA01ProcessamentoDePasta:
    def test_lista_classifica_e_copia_os_arquivos(self, full_company):
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)
        _write(full_company, "pix.txt", PIX)

        batch = process_company_folder(full_company)

        assert batch.total == 2
        assert len(batch.errors) == 0
        for result in batch.results:
            assert result.destination_path is not None
            from pathlib import Path

            assert Path(result.destination_path).exists()

    def test_arquivo_original_permanece_na_entrada(self, full_company):
        source = _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)

        process_company_folder(full_company)

        assert source.exists(), "RNF02: os arquivos são copiados, nunca movidos"

    def test_estrutura_ano_mes_tipo_e_criada(self, full_company):
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)

        batch = process_company_folder(full_company)
        destination = batch.results[0].destination_path

        assert "2026" in destination
        assert "01 - Janeiro" in destination
        assert "CONTAS_PAGAS" in destination


class TestCA02TratamentoDeErro:
    def test_extensao_nao_suportada_vai_para_erros(self, full_company):
        _write(full_company, "arquivo.exe", "binário qualquer")

        batch = process_company_folder(full_company)

        assert len(batch.errors) == 1
        assert "ERROS" in batch.errors[0].destination_path

    def test_erro_grava_marcador_json_com_motivo(self, full_company):
        import json
        from pathlib import Path

        _write(full_company, "arquivo.exe", "x")

        batch = process_company_folder(full_company)
        marker = Path(batch.errors[0].destination_path + ".erro.json")

        assert marker.exists()
        payload = json.loads(marker.read_text(encoding="utf-8"))
        assert payload["erro_tipo"]
        assert payload["acao_recomendada"]

    def test_erro_fica_registrado_no_banco(self, full_company):
        _write(full_company, "arquivo.exe", "x")

        batch = process_company_folder(full_company)
        document = document_repository.get_document(batch.errors[0].document_id)

        assert document.processing_status == DocumentStatus.ERROR
        assert document.error_message

    def test_um_arquivo_com_erro_nao_interrompe_o_lote(self, full_company):
        _write(full_company, "arquivo.exe", "x")
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)

        batch = process_company_folder(full_company)

        assert batch.total == 2
        assert len(batch.errors) == 1
        assert len(batch.classified) == 1


class TestCA03DocumentoDesconhecido:
    def test_texto_irreconhecivel_vai_para_outros(self, full_company):
        _write(full_company, "misterio.txt", ILEGIVEL)

        batch = process_company_folder(full_company)

        assert len(batch.unknown) == 1
        assert "OUTROS" in batch.unknown[0].destination_path

    def test_documento_desconhecido_nao_gera_lancamento(self, full_company):
        _write(full_company, "misterio.txt", ILEGIVEL)

        batch = process_company_folder(full_company)
        document = document_repository.get_document(batch.unknown[0].document_id)
        result = generate_entry_for_document(document, full_company)

        assert result.entry is None
        assert document.document_type == DocumentType.OUTRO


class TestCA04Duplicidade:
    def test_mesmo_arquivo_reprocessado_vira_duplicado(self, full_company):
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)

        process_company_folder(full_company)
        second = process_company_folder(full_company)

        assert len(second.duplicates) == 1
        assert "DUPLICADOS" in second.duplicates[0].destination_path

    def test_duplicado_nao_gera_lancamento(self, full_company):
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)
        process_company_folder(full_company)
        second = process_company_folder(full_company)

        document = document_repository.get_document(second.duplicates[0].document_id)
        result = generate_entry_for_document(document, full_company)

        assert result.entry is None
        assert "DUPLICATE" in (result.skipped_reason or "")

    def test_mesmo_fato_em_arquivo_diferente_e_detectado(self, full_company):
        """Mesmo credor, valor, data e documento, em arquivo de outro nome e conteúdo.

        O byte extra muda o hash, então a segunda camada — dados essenciais —
        é a única que pode pegar esta duplicidade.
        """
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)
        process_company_folder(full_company)

        copy_path = _write(full_company, "copia_contas.txt", CONTAS_PAGAS + "\n")
        second = process_company_folder(full_company)

        by_business_data = next(
            r for r in second.duplicates if r.file_name == copy_path.name
        )
        assert "Mesma data, valor e documento" in by_business_data.message


class TestHierarquiaDeDocumentoBase:
    def test_comprovante_pix_nao_gera_lancamento_quando_ha_contas_pagas(self, full_company):
        """PRD seção 10: contas pagas é o documento base; o PIX vira evidência."""
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)
        _write(full_company, "pix.txt", PIX)
        process_company_folder(full_company)

        documents = document_repository.search_documents(company_id=full_company.id)
        contas = next(d for d in documents if d.document_type == DocumentType.CONTAS_PAGAS)
        pix = next(d for d in documents if d.document_type == DocumentType.COMPROVANTE_PIX)

        contas_result = generate_entry_for_document(contas, full_company)
        pix_result = generate_entry_for_document(pix, full_company)

        assert contas_result.entry is not None, "contas pagas deve gerar o lançamento"
        assert pix_result.entry is None, "o PIX do mesmo fato não pode gerar lançamento"
        assert "maior riqueza contábil" in (pix_result.skipped_reason or "")

    def test_pix_isolado_gera_lancamento(self, full_company):
        _write(full_company, "pix.txt", PIX)
        process_company_folder(full_company)

        document = document_repository.search_documents(company_id=full_company.id)[0]
        result = generate_entry_for_document(document, full_company)

        assert result.entry is not None


class TestCA05ExemplosNoContexto:
    def test_exemplo_relevante_entra_no_prompt(self, full_company):
        from core.accounting.prompt_builder import build_prompt
        from db.repositories import example_repository

        example_repository.create_example(
            full_company.id,
            {
                "title": "Pagamento de prestador por contas pagas",
                "document_type": "CONTAS_PAGAS",
                "base_document": "Relatório de contas pagas",
                "counterparty": "FORNECEDOR FICTICIO SERVICOS LTDA",
                "keywords": "empreiteiro; construção",
                "debit_account_code": "3.1.01.001",
                "credit_account_code": "1.1.01.001",
                "history_complement_template": "Pagamento a {fornecedor}",
                "accounting_explanation": "Contas pagas tem credor, documento e líquido.",
                "priority": "Alta",
                "active": True,
            },
        )

        context = build_prompt(
            company=full_company,
            document_type="CONTAS_PAGAS",
            extracted_data={"counterparty": "FORNECEDOR FICTICIO SERVICOS LTDA"},
            document_text=CONTAS_PAGAS,
            counterparty="FORNECEDOR FICTICIO SERVICOS LTDA",
        )

        assert len(context.examples_used) == 1
        assert "Pagamento de prestador por contas pagas" in context.user_prompt

    def test_exemplo_de_outro_tipo_e_contraparte_nao_entra(self, full_company):
        from core.accounting.prompt_builder import build_prompt
        from db.repositories import example_repository

        example_repository.create_example(
            full_company.id,
            {
                "title": "Receita de serviços",
                "document_type": "FATURAMENTO",
                "base_document": "Relatório de faturamento",
                "counterparty": "OUTRA EMPRESA SA",
                "debit_account_code": "1.1.01.001",
                "credit_account_code": "3.1.01.001",
                "history_complement_template": "Receita de {cliente}",
                "accounting_explanation": "Faturamento gera receita.",
                "priority": "Baixa",
                "active": True,
            },
        )

        context = build_prompt(
            company=full_company,
            document_type="CONTAS_PAGAS",
            extracted_data={},
            document_text=CONTAS_PAGAS,
            counterparty="FORNECEDOR FICTICIO SERVICOS LTDA",
        )

        assert context.examples_used == []

    def test_prompt_proibe_inventar_contas(self, full_company):
        from core.accounting.prompt_builder import build_prompt

        context = build_prompt(
            company=full_company,
            document_type="CONTAS_PAGAS",
            extracted_data={},
            document_text=CONTAS_PAGAS,
        )

        assert "NÃO invente contas" in context.user_prompt
        assert "3.1.01.001" in context.user_prompt
        assert "3.1.01 |" not in context.user_prompt, "conta sintética não deve ser oferecida"


class TestCA06ValidacaoHumana:
    def test_lancamento_nasce_pendente_de_revisao(self, full_company):
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)
        process_company_folder(full_company)

        document = document_repository.search_documents(company_id=full_company.id)[0]
        result = generate_entry_for_document(document, full_company)

        assert result.entry.status == EntryStatus.PENDING_REVIEW

    def test_aprovacao_com_correcao_registra_memoria(self, full_company):
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)
        process_company_folder(full_company)
        document = document_repository.search_documents(company_id=full_company.id)[0]
        entry = generate_entry_for_document(document, full_company).entry

        original = accounting_entry_repository.entry_to_dict(entry)
        corrected = {**original, "history_complement": "Histórico corrigido pelo contador"}

        accounting_entry_repository.update_entry(
            entry.id, {"history_complement": corrected["history_complement"]}
        )
        accounting_entry_repository.approve_entry(entry.id, "contador", corrected=True)
        accounting_entry_repository.create_correction(
            entry_id=entry.id,
            company_id=full_company.id,
            document_id=document.id,
            document_type=document.document_type,
            counterparty=document.counterparty,
            original=original,
            corrected=corrected,
            corrective_guidance="Sempre citar o número do documento no histórico.",
            convert_to_example=True,
            created_by="contador",
        )

        updated = accounting_entry_repository.get_entry(entry.id)
        corrections = accounting_entry_repository.list_corrections(full_company.id)

        assert updated.status == EntryStatus.CORRECTED_APPROVED
        assert updated.approved_by == "contador"
        assert len(corrections) == 1
        assert corrections[0].corrective_guidance

    def test_correcao_vira_exemplo_apos_confirmacao(self, full_company):
        from core.accounting import correction_memory
        from db.repositories import example_repository

        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)
        process_company_folder(full_company)
        document = document_repository.search_documents(company_id=full_company.id)[0]
        entry = generate_entry_for_document(document, full_company).entry

        correction = accounting_entry_repository.create_correction(
            entry_id=entry.id,
            company_id=full_company.id,
            document_id=document.id,
            document_type=document.document_type,
            counterparty=document.counterparty,
            original=accounting_entry_repository.entry_to_dict(entry),
            corrected={
                "debit_account_code": "3.1.01.001",
                "credit_account_code": "1.1.01.001",
                "history_complement": "Pagamento a fornecedor",
            },
            corrective_guidance="Usar serviços de terceiros para este fornecedor.",
            convert_to_example=True,
            created_by="contador",
        )

        example_id = correction_memory.convert_correction_to_example(correction)

        assert example_repository.get_example(example_id) is not None
        assert accounting_entry_repository.list_corrections(full_company.id)[0].converted_example_id == example_id


class TestCA07ExportacaoPontaAPonta:
    def test_fluxo_completo_ate_o_xlsx(self, full_company):
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)
        process_company_folder(full_company)

        document = document_repository.search_documents(company_id=full_company.id)[0]
        entry = generate_entry_for_document(document, full_company).entry
        accounting_entry_repository.approve_entry(entry.id, "contador")

        result = export_entries(full_company, year=2026, month=1, exported_by="contador")

        assert result.entry_count == 1
        assert result.file_path.exists()

        sheet = load_workbook(result.file_path).active
        assert sheet.max_row == 2
        assert sheet.cell(row=2, column=1).value == "06/01/2026"

    def test_documento_com_erro_nunca_chega_ao_xlsx(self, full_company):
        import pytest

        from core.exceptions import ExportBlockedError

        _write(full_company, "arquivo.exe", "x")
        process_company_folder(full_company)

        with pytest.raises(ExportBlockedError):
            export_entries(full_company, year=2026, month=1, exported_by="contador")


class TestCA08Consulta:
    def test_consulta_por_empresa_e_competencia(self, full_company):
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)
        process_company_folder(full_company)

        found = document_repository.search_documents(
            company_id=full_company.id, year=2026, month=1
        )

        assert len(found) == 1
        assert found[0].original_file_name == "contas_pagas.txt"

    def test_rastreabilidade_arquivo_documento_lancamento_lote(self, full_company):
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)
        process_company_folder(full_company)

        document = document_repository.search_documents(company_id=full_company.id)[0]
        entry = generate_entry_for_document(document, full_company).entry
        accounting_entry_repository.approve_entry(entry.id, "contador")
        result = export_entries(full_company, year=2026, month=1, exported_by="contador")

        final_entry = accounting_entry_repository.get_entry(entry.id)

        assert final_entry.document_id == document.id
        assert final_entry.export_batch_id == result.batch_id
        assert document.original_file_path
        assert document.stored_file_path

    def test_contadores_do_dashboard(self, full_company):
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)
        _write(full_company, "misterio.txt", ILEGIVEL)
        process_company_folder(full_company)

        counts = document_repository.count_documents_by_status(full_company.id)

        assert counts.get(DocumentStatus.EXTRACTED) == 1
        assert counts.get(DocumentStatus.UNKNOWN) == 1


class TestScanner:
    def test_ignora_arquivos_temporarios(self, full_company):
        _write(full_company, "~$temp.xlsx", "x")
        _write(full_company, "contas_pagas.txt", CONTAS_PAGAS)

        files = scan_folder(full_company.input_folder_path)

        assert [f.name for f in files] == ["contas_pagas.txt"]

    def test_marca_extensao_nao_suportada(self, full_company):
        _write(full_company, "arquivo.exe", "x")

        files = scan_folder(full_company.input_folder_path)

        assert files[0].supported is False

    def test_pasta_inexistente_retorna_lista_vazia(self):
        assert scan_folder("/caminho/que/nao/existe") == []
