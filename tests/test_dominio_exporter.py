"""Testes do exportador Domínio e da trava de lançamentos não aprovados."""

from __future__ import annotations

from datetime import date

import pytest

from core.constants import DOMINIO_COLUMNS, EntryStatus
from core.dominio.dominio_exporter import export_entries, preview_export
from core.dominio.dominio_template import entry_to_row, format_amount, format_date
from core.dominio.export_validator import validate_entry_for_export, validate_export_batch
from core.exceptions import ExportBlockedError
from db.repositories import accounting_entry_repository, document_repository, export_repository


def make_entry(company_id: int, *, status=EntryStatus.APPROVED, **overrides):
    """Documento + lançamento prontos para exportação."""
    document = document_repository.create_document(
        {
            "company_id": company_id,
            "original_file_name": "contas_pagas.pdf",
            "original_file_path": "/x/contas_pagas.pdf",
            "file_hash": overrides.pop("file_hash", "e" * 64),
            "file_extension": ".pdf",
            "document_type": "CONTAS_PAGAS",
            "competence_year": 2026,
            "competence_month": 1,
            "processing_status": "ENTRY_GENERATED",
        }
    )
    data = {
        "company_id": company_id,
        "document_id": document.id,
        "entry_date": date(2026, 1, 6),
        "debit_account_code": "3.1.01.001",
        "credit_account_code": "1.1.01.001",
        "amount": 1500.00,
        "history_code": "101",
        "history_complement": "Pagamento a FICTICIO conforme NF-2026-0042",
        "branch_code": "1",
        "status": status,
    }
    data.update(overrides)
    return accounting_entry_repository.create_entry(data)


class TestDominioTemplate:
    def test_data_no_formato_brasileiro(self):
        assert format_date(date(2026, 1, 6)) == "06/01/2026"
        assert format_date(None) == ""

    def test_valor_numerico_com_duas_casas(self):
        assert format_amount(1500) == 1500.00
        assert isinstance(format_amount(1500), float)
        assert format_amount(None) == 0.0

    def test_linha_tem_todas_as_colunas_do_layout(self, full_company):
        entry = make_entry(full_company.id)
        row = entry_to_row(entry, is_first=True)

        assert set(row) == set(DOMINIO_COLUMNS)
        assert row["Inicia Lote"] == "1"
        assert row["Cód. Conta Debito"] == "3.1.01.001"
        assert row["Valor"] == 1500.00

    def test_apenas_a_primeira_linha_inicia_lote(self, full_company):
        first = entry_to_row(make_entry(full_company.id), is_first=True)
        second = entry_to_row(make_entry(full_company.id, file_hash="f" * 64), is_first=False)

        assert first["Inicia Lote"] == "1"
        assert second["Inicia Lote"] == ""


class TestExportValidator:
    def test_lancamento_aprovado_e_completo_passa(self, full_company):
        assert validate_entry_for_export(make_entry(full_company.id)) == []

    @pytest.mark.parametrize(
        "status",
        [
            EntryStatus.PENDING_REVIEW,
            EntryStatus.REJECTED,
            EntryStatus.DUPLICATE,
            EntryStatus.ERROR,
        ],
    )
    def test_status_nao_aprovado_e_bloqueado(self, full_company, status):
        reasons = validate_entry_for_export(make_entry(full_company.id, status=status))

        assert reasons
        assert any("não é exportável" in reason for reason in reasons)

    def test_status_corrigido_e_aprovado_passa(self, full_company):
        entry = make_entry(full_company.id, status=EntryStatus.CORRECTED_APPROVED)
        assert validate_entry_for_export(entry) == []

    @pytest.mark.parametrize(
        "field,value,expected",
        [
            ("entry_date", None, "Data"),
            ("debit_account_code", None, "débito"),
            ("credit_account_code", None, "crédito"),
            ("amount", None, "Valor"),
            ("history_complement", "", "histórico"),
        ],
    )
    def test_campo_obrigatorio_ausente_bloqueia(self, full_company, field, value, expected):
        entry = make_entry(full_company.id, **{field: value})
        reasons = validate_entry_for_export(entry)

        assert any(expected in reason for reason in reasons)

    def test_conta_desativada_apos_aprovacao_bloqueia_o_lote(self, full_company):
        from db.repositories import company_repository

        entry = make_entry(full_company.id)
        company_repository.upsert_account(
            full_company.id,
            {
                "account_code": "3.1.01.001",
                "description": "Serviços de terceiros",
                "accepts_entries": True,
                "active": False,
            },
        )

        result = validate_export_batch([entry], company_id=full_company.id)

        assert not result.is_valid
        assert any(
            "desativada" in reason for issue in result.blocked for reason in issue.reasons
        )

    def test_lote_vazio_gera_aviso(self, full_company):
        result = validate_export_batch([], company_id=full_company.id)

        assert not result.is_valid
        assert result.warnings


class TestPreviewExport:
    def test_previa_traz_somente_aprovados(self, full_company):
        make_entry(full_company.id, status=EntryStatus.APPROVED)
        make_entry(full_company.id, status=EntryStatus.PENDING_REVIEW, file_hash="1" * 64)
        make_entry(full_company.id, status=EntryStatus.REJECTED, file_hash="2" * 64)

        rows, validation = preview_export(full_company.id, year=2026, month=1)

        assert len(rows) == 1
        assert len(validation.valid_entries) == 1

    def test_previa_nao_altera_status(self, full_company):
        entry = make_entry(full_company.id)

        preview_export(full_company.id, year=2026, month=1)

        assert accounting_entry_repository.get_entry(entry.id).status == EntryStatus.APPROVED


class TestExportEntries:
    def test_gera_xlsx_com_cabecalho_correto(self, full_company):
        from openpyxl import load_workbook

        make_entry(full_company.id)

        result = export_entries(full_company, year=2026, month=1, exported_by="contador")

        assert result.file_path.exists()
        sheet = load_workbook(result.file_path).active
        header = [cell.value for cell in sheet[1]]
        assert header == list(DOMINIO_COLUMNS)
        assert sheet.max_row == 2

    def test_marca_lancamentos_como_exportados(self, full_company):
        entry = make_entry(full_company.id)

        result = export_entries(full_company, year=2026, month=1, exported_by="contador")

        updated = accounting_entry_repository.get_entry(entry.id)
        assert updated.status == EntryStatus.EXPORTED
        assert updated.export_batch_id == result.batch_id

    def test_registra_lote_com_totais(self, full_company):
        make_entry(full_company.id)
        make_entry(full_company.id, amount=500.00, file_hash="3" * 64)

        result = export_entries(full_company, year=2026, month=1, exported_by="contador")
        batch = export_repository.get_batch(result.batch_id)

        assert batch.entry_count == 2
        assert batch.total_amount == 2000.00
        assert batch.exported_by == "contador"

    def test_lancamento_pendente_nunca_e_exportado(self, full_company):
        """RNF03 — a trava central da exportação."""
        pending = make_entry(full_company.id, status=EntryStatus.PENDING_REVIEW)
        make_entry(full_company.id, status=EntryStatus.APPROVED, file_hash="4" * 64)

        result = export_entries(full_company, year=2026, month=1, exported_by="contador")

        assert result.entry_count == 1
        assert accounting_entry_repository.get_entry(pending.id).status == EntryStatus.PENDING_REVIEW

    def test_sem_aprovados_levanta_erro_de_bloqueio(self, full_company):
        make_entry(full_company.id, status=EntryStatus.PENDING_REVIEW)

        with pytest.raises(ExportBlockedError):
            export_entries(full_company, year=2026, month=1, exported_by="contador")

    def test_lancamento_incompleto_bloqueia_o_lote_inteiro(self, full_company):
        make_entry(full_company.id)
        make_entry(full_company.id, amount=None, file_hash="5" * 64)

        with pytest.raises(ExportBlockedError):
            export_entries(full_company, year=2026, month=1, exported_by="contador")

    def test_exportacao_parcial_quando_autorizada(self, full_company):
        make_entry(full_company.id)
        make_entry(full_company.id, amount=None, file_hash="6" * 64)

        result = export_entries(
            full_company, year=2026, month=1, exported_by="contador", allow_partial=True
        )

        assert result.entry_count == 1

    def test_lote_aparece_no_historico_com_seus_lancamentos(self, full_company):
        entry = make_entry(full_company.id)

        result = export_entries(full_company, year=2026, month=1, exported_by="contador")
        entries = export_repository.list_batch_entries(result.batch_id)

        assert [e.id for e in entries] == [entry.id]

    def test_reexportacao_nao_duplica_lancamento_ja_exportado(self, full_company):
        make_entry(full_company.id)
        export_entries(full_company, year=2026, month=1, exported_by="contador")

        with pytest.raises(ExportBlockedError):
            export_entries(full_company, year=2026, month=1, exported_by="contador")
