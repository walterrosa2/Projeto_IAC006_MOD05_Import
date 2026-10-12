"""Geração do XLSX no layout do Domínio Contábil (Tarefa 7.1)."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from core.config import get_settings
from core.constants import DOMINIO_COLUMNS
from core.dominio.dominio_template import entries_to_rows
from core.dominio.export_validator import ExportValidationResult, validate_export_batch
from core.exceptions import ExportBlockedError
from core.logging_config import get_logger, log_audit
from db.models import AccountingEntry, Company
from db.repositories import accounting_entry_repository, export_repository

logger = get_logger()


@dataclass
class ExportResult:
    """Resultado da exportação concluída."""

    file_path: Path
    batch_id: int
    entry_count: int
    total_amount: float
    validation: ExportValidationResult


def _write_workbook(rows: list[dict], destination: Path) -> None:
    """Escreve o XLSX com cabeçalho e formatação mínima."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lancamentos"

    sheet.append(list(DOMINIO_COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append([row.get(column, "") for column in DOMINIO_COLUMNS])

    # Valor como número com 2 casas; o Domínio rejeita valor em texto.
    value_column = DOMINIO_COLUMNS.index("Valor") + 1
    for row_index in range(2, sheet.max_row + 1):
        sheet.cell(row=row_index, column=value_column).number_format = "0.00"

    for index, column in enumerate(DOMINIO_COLUMNS, start=1):
        width = max(len(column) + 4, 14)
        if column == "Complemento Histórico":
            width = 50
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)


def preview_export(
    company_id: int, *, year: int | None = None, month: int | None = None
) -> tuple[list[dict], ExportValidationResult]:
    """Prévia da exportação, sem gravar arquivo nem alterar status.

    Returns:
        (linhas no layout Domínio, resultado da validação).
    """
    entries = accounting_entry_repository.list_exportable_entries(company_id, year, month)
    validation = validate_export_batch(entries, company_id=company_id)
    return entries_to_rows(validation.valid_entries), validation


def export_entries(
    company: Company,
    *,
    year: int | None = None,
    month: int | None = None,
    exported_by: str | None = None,
    entries: list[AccountingEntry] | None = None,
    allow_partial: bool = False,
) -> ExportResult:
    """Gera o XLSX e registra o lote, marcando os lançamentos como EXPORTED.

    Args:
        entries: subconjunto escolhido na tela; None busca todos os aprovados.
        allow_partial: exporta os liberados mesmo havendo bloqueados. Por padrão
            um único bloqueio impede o lote inteiro (RNF03).

    Raises:
        ExportBlockedError: nada a exportar ou bloqueios com `allow_partial=False`.
    """
    candidates = (
        entries
        if entries is not None
        else accounting_entry_repository.list_exportable_entries(company.id, year, month)
    )
    validation = validate_export_batch(candidates, company_id=company.id)

    if validation.blocked and not allow_partial:
        detail = "; ".join(
            f"#{issue.entry_id}: {', '.join(issue.reasons)}" for issue in validation.blocked[:5]
        )
        log_audit(
            "export_blocked",
            actor="user",
            target=f"empresa:{company.id}",
            outcome="blocked",
            data={"bloqueados": validation.blocked_count, "detalhe": detail},
        )
        raise ExportBlockedError(
            f"Exportação bloqueada: {validation.blocked_count} lançamento(s) com pendência.",
            detail=detail,
        )

    if not validation.valid_entries:
        raise ExportBlockedError(
            "Nenhum lançamento aprovado disponível para exportação.",
            detail="Aprove lançamentos na tela de validação antes de exportar.",
        )

    settings = get_settings()
    file_name = export_repository.build_export_filename(company.name, year, month)
    destination = settings.export_dir / file_name

    rows = entries_to_rows(validation.valid_entries)
    _write_workbook(rows, destination)

    total_amount = round(sum(e.amount or 0.0 for e in validation.valid_entries), 2)
    batch = export_repository.create_batch(
        company_id=company.id,
        period_year=year,
        period_month=month,
        file_path=str(destination),
        entry_ids=[e.id for e in validation.valid_entries],
        total_amount=total_amount,
        exported_by=exported_by,
    )

    log_audit(
        "export_completed",
        actor="user",
        target=f"lote:{batch.id}",
        outcome="success",
        data={
            "empresa": company.name,
            "competencia": f"{year}-{month}" if year and month else "completo",
            "lancamentos": len(validation.valid_entries),
            "valor_total": total_amount,
            "arquivo": str(destination),
            "bloqueados_ignorados": validation.blocked_count if allow_partial else 0,
        },
    )
    logger.info(
        f"Exportação concluída: lote {batch.id}, "
        f"{len(validation.valid_entries)} lançamento(s), arquivo {destination.name}"
    )

    return ExportResult(
        file_path=destination,
        batch_id=batch.id,
        entry_count=len(validation.valid_entries),
        total_amount=total_amount,
        validation=validation,
    )
